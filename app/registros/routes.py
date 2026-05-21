from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
from io import BytesIO
import mimetypes
import os
import re
import unicodedata
import uuid
import zipfile

from openpyxl import load_workbook
from werkzeug.utils import secure_filename

from flask import (
    Blueprint,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    current_app,
    send_file,
    abort
)
from flask_login import login_required, current_user

from app.extensions import db
from app.models import Entry, Supplier, BudgetItem, EntryDocument
from app.services.budget_alert_service import (
    refresh_alerts_after_entry_change,
    refresh_alerts_for_budget_item_ids,
)


registros_bp = Blueprint("registros", __name__, url_prefix="/registros")


MESES_ES = {
    1: "ENERO",
    2: "FEBRERO",
    3: "MARZO",
    4: "ABRIL",
    5: "MAYO",
    6: "JUNIO",
    7: "JULIO",
    8: "AGOSTO",
    9: "SEPTIEMBRE",
    10: "OCTUBRE",
    11: "NOVIEMBRE",
    12: "DICIEMBRE",
}


DOCUMENT_LABELS = {
    "payment_request": "Solicitud de pago",
    "payment_receipt": "Comprobante de pago",
    "gmail_comments": "Comentarios vía Gmail",
}


DOCUMENT_FIELD_MAP = {
    "payment_request_file": "payment_request",
    "payment_receipt_file": "payment_receipt",
    "gmail_comments_file": "gmail_comments",
}


ALLOWED_DOCUMENT_EXTENSIONS = {
    "pdf",
    "png",
    "jpg",
    "jpeg",
    "xlsx",
    "xlsm",
    "xls",
}


PREVIEWABLE_EXTENSIONS = {
    "pdf",
    "png",
    "jpg",
    "jpeg",
}


HEADER_MAP = {
    "FECHA": "fecha",
    "MONTO": "monto",
    "PROVEEDOR": "proveedor",
    "RUBRO": "rubro",
    "FOLIO DE OPERACION": "folio",
    "CONCEPTO": "concepto",
    "COMENTARIOS": "comentarios",
}


REQUIRED_IMPORT_FIELDS = [
    "fecha",
    "monto",
    "proveedor",
    "rubro",
    "folio",
    "concepto",
    "comentarios",
]


LEGAL_WORDS = {
    "S", "A", "C", "V", "SA", "CV", "SACV", "SAPI", "DE", "DEL", "LA", "LAS", "LOS",
    "SOCIEDAD", "ANONIMA", "CAPITAL", "VARIABLE", "RL", "MI", "SC", "SPR",
}


# =========================================================
# PARSERS
# =========================================================

def parse_date(value):
    if not value:
        return None

    value = str(value).strip()

    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue

    return None


def parse_amount(value):
    if value is None:
        return None

    value = str(value).strip()

    if not value:
        return None

    cleaned = (
        value.replace("$", "")
        .replace(",", "")
        .replace("MXN", "")
        .replace(" ", "")
        .strip()
    )

    try:
        return Decimal(cleaned)
    except InvalidOperation:
        return None


def parse_excel_date(value):
    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    return parse_date(str(value).strip())


def parse_excel_amount(value):
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))

    return parse_amount(str(value).strip() if value is not None else "")


# =========================================================
# NORMALIZACIÓN / MATCHING
# =========================================================

def normalize_text(value):
    if value is None:
        return ""

    text = str(value).strip().upper()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Z0-9.\s]", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_header(value):
    text = normalize_text(value)
    text = text.replace(".", " ")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def normalize_supplier_name(value):
    text = normalize_text(value)
    text = text.replace("S.A. DE C.V.", " ")
    text = text.replace("S A DE C V", " ")
    text = text.replace("SA DE CV", " ")
    text = text.replace("S DE RL DE CV", " ")
    text = text.replace("SAPI DE CV", " ")

    tokens = [
        token
        for token in text.replace(".", " ").split()
        if token not in LEGAL_WORDS
    ]

    return " ".join(tokens)


def token_sort_text(value):
    return " ".join(sorted(normalize_supplier_name(value).split()))


def similarity_score(a, b):
    a_norm = normalize_supplier_name(a)
    b_norm = normalize_supplier_name(b)

    if not a_norm or not b_norm:
        return 0

    direct = SequenceMatcher(None, a_norm, b_norm).ratio()
    token_sort = SequenceMatcher(None, token_sort_text(a_norm), token_sort_text(b_norm)).ratio()

    a_tokens = set(a_norm.split())
    b_tokens = set(b_norm.split())
    token_overlap = len(a_tokens & b_tokens) / max(len(a_tokens), len(b_tokens), 1)

    return max(direct, token_sort, token_overlap)


def find_supplier_match(raw_supplier, suppliers, min_score=0.78):
    best_supplier = None
    best_score = 0

    for supplier in suppliers:
        candidates = [supplier.business_name]

        if getattr(supplier, "rfc", None):
            candidates.append(supplier.rfc)

        for candidate in candidates:
            score = similarity_score(raw_supplier, candidate)

            if normalize_text(raw_supplier) == normalize_text(candidate):
                score = 1

            if score > best_score:
                best_score = score
                best_supplier = supplier

    if best_supplier and best_score >= min_score:
        return best_supplier, best_score

    return None, best_score


def find_budget_item_match(raw_rubro, budget_items):
    rubro_text = normalize_text(raw_rubro).replace(" ", "")

    for item in budget_items:
        item_code = normalize_text(item.item_code).replace(" ", "")

        if rubro_text == item_code:
            return item

    return None


# =========================================================
# IMPORTACIÓN EXCEL
# =========================================================

def get_import_headers(ws):
    headers = {}

    for col in range(1, ws.max_column + 1):
        raw_header = ws.cell(row=1, column=col).value
        normalized = normalize_header(raw_header)

        if normalized in HEADER_MAP:
            headers[HEADER_MAP[normalized]] = col

    return headers


def generate_unique_key_for_import(application_year, counters):
    if application_year not in counters:
        prefix = f"APL-{application_year}-"

        last_entry = (
            Entry.query
            .filter(Entry.unique_key.like(f"{prefix}%"))
            .order_by(Entry.id.desc())
            .first()
        )

        if not last_entry:
            counters[application_year] = 0
        else:
            try:
                counters[application_year] = int(last_entry.unique_key.split("-")[-1])
            except (ValueError, IndexError):
                counters[application_year] = 0

    counters[application_year] += 1
    return f"APL-{application_year}-{counters[application_year]:06d}"


# =========================================================
# RUBROS / CLAVES
# =========================================================

def natural_code_key(item_code):
    try:
        return tuple(int(part) for part in item_code.split("."))
    except ValueError:
        return tuple(item_code.split("."))


def get_sorted_budget_items():
    return sorted(
        BudgetItem.query.filter_by(is_active=True).all(),
        key=lambda x: natural_code_key(x.item_code)
    )


def generate_unique_key(application_year):
    prefix = f"APL-{application_year}-"

    last_entry = (
        Entry.query
        .filter(Entry.unique_key.like(f"{prefix}%"))
        .order_by(Entry.id.desc())
        .first()
    )

    if not last_entry:
        next_number = 1
    else:
        try:
            last_number = int(last_entry.unique_key.split("-")[-1])
            next_number = last_number + 1
        except (ValueError, IndexError):
            next_number = 1

    return f"{prefix}{next_number:06d}"


# =========================================================
# DOCUMENTOS / EVIDENCIAS
# =========================================================

def allowed_document(filename):
    if not filename or "." not in filename:
        return False

    extension = filename.rsplit(".", 1)[1].lower()
    return extension in ALLOWED_DOCUMENT_EXTENSIONS


def get_upload_base_dir():
    return os.path.join(
        current_app.root_path,
        "static",
        "uploads",
        "entry_documents"
    )


def ensure_directory(path):
    os.makedirs(path, exist_ok=True)


def save_entry_document(entry, uploaded_file, document_type):
    if not uploaded_file or not uploaded_file.filename:
        return None

    if not allowed_document(uploaded_file.filename):
        raise ValueError("Formato no permitido. Usa PDF, Excel, PNG o JPG.")

    original_filename = secure_filename(uploaded_file.filename)
    extension = original_filename.rsplit(".", 1)[1].lower()

    year_folder = str(entry.application_year)
    entry_folder = secure_filename(entry.unique_key)

    destination_dir = os.path.join(
        get_upload_base_dir(),
        year_folder,
        entry_folder
    )

    ensure_directory(destination_dir)

    stored_filename = f"{document_type}_{uuid.uuid4().hex}.{extension}"
    absolute_path = os.path.join(destination_dir, stored_filename)

    uploaded_file.save(absolute_path)

    size_kb = round(os.path.getsize(absolute_path) / 1024, 2)
    mime_type = mimetypes.guess_type(absolute_path)[0] or "application/octet-stream"

    previous_document = (
        EntryDocument.query
        .filter_by(entry_id=entry.id, document_type=document_type)
        .first()
    )

    if previous_document:
        try:
            if previous_document.file_path and os.path.exists(previous_document.file_path):
                os.remove(previous_document.file_path)
        except OSError:
            pass

        previous_document.original_filename = original_filename
        previous_document.stored_filename = stored_filename
        previous_document.file_path = absolute_path
        previous_document.file_extension = extension
        previous_document.mime_type = mime_type
        previous_document.file_size_kb = size_kb
        previous_document.uploaded_by = current_user.id
        previous_document.uploaded_at = datetime.utcnow()

        return previous_document

    document = EntryDocument(
        entry_id=entry.id,
        document_type=document_type,
        original_filename=original_filename,
        stored_filename=stored_filename,
        file_path=absolute_path,
        file_extension=extension,
        mime_type=mime_type,
        file_size_kb=size_kb,
        uploaded_by=current_user.id
    )

    db.session.add(document)
    return document


def save_entry_documents_from_request(entry):
    for field_name, document_type in DOCUMENT_FIELD_MAP.items():
        uploaded_file = request.files.get(field_name)

        if uploaded_file and uploaded_file.filename:
            save_entry_document(entry, uploaded_file, document_type)


def get_document_map(entry_id):
    documents = EntryDocument.query.filter_by(entry_id=entry_id).all()
    return {document.document_type: document for document in documents}


def remove_document_file(document):
    try:
        if document.file_path and os.path.exists(document.file_path):
            os.remove(document.file_path)
    except OSError:
        pass


# =========================================================
# FILTROS
# =========================================================

def build_filtered_query():
    query = (
        Entry.query
        .join(Supplier, Entry.supplier_id == Supplier.id)
        .outerjoin(BudgetItem, Entry.budget_item_id == BudgetItem.id)
    )

    search = request.args.get("q", "").strip()
    fecha_desde = request.args.get("fecha_desde", "").strip()
    fecha_hasta = request.args.get("fecha_hasta", "").strip()
    mes = request.args.get("mes", "").strip()
    supplier_id = request.args.get("supplier_id", "").strip()
    budget_item_id = request.args.get("budget_item_id", "").strip()
    folio = request.args.get("folio", "").strip()
    sort = request.args.get("sort", "fecha").strip()
    direction = request.args.get("direction", "desc").strip()

    if search:
        like = f"%{search}%"

        query = query.filter(
            db.or_(
                Entry.unique_key.like(like),
                Supplier.business_name.like(like),
                Supplier.rfc.like(like),
                Entry.operation_folio.like(like),
                Entry.concept.like(like),
                Entry.comments.like(like),
                BudgetItem.item_code.like(like),
                BudgetItem.item_name.like(like),
            )
        )

    parsed_fecha_desde = parse_date(fecha_desde)
    parsed_fecha_hasta = parse_date(fecha_hasta)

    if parsed_fecha_desde:
        query = query.filter(Entry.application_date >= parsed_fecha_desde)

    if parsed_fecha_hasta:
        query = query.filter(Entry.application_date <= parsed_fecha_hasta)

    if mes:
        try:
            query = query.filter(Entry.application_month_number == int(mes))
        except ValueError:
            pass

    if supplier_id:
        try:
            query = query.filter(Entry.supplier_id == int(supplier_id))
        except ValueError:
            pass

    if budget_item_id:
        try:
            query = query.filter(Entry.budget_item_id == int(budget_item_id))
        except ValueError:
            pass

    if folio:
        query = query.filter(Entry.operation_folio.like(f"%{folio}%"))

    sort_options = {
        "fecha": Entry.application_date,
        "monto": Entry.amount,
        "proveedor": Supplier.business_name,
        "rubro": BudgetItem.item_code,
        "folio": Entry.operation_folio,
        "clave": Entry.unique_key,
        "creado": Entry.created_at,
    }

    sort_column = sort_options.get(sort, Entry.application_date)

    if direction == "asc":
        query = query.order_by(sort_column.asc(), Entry.id.asc())
    else:
        query = query.order_by(sort_column.desc(), Entry.id.desc())

    filters = {
        "q": search,
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "mes": mes,
        "supplier_id": supplier_id,
        "budget_item_id": budget_item_id,
        "folio": folio,
        "sort": sort,
        "direction": direction,
    }

    return query, filters


# =========================================================
# RUTAS
# =========================================================

@registros_bp.route("/")
@login_required
def index():
    page = request.args.get("page", 1, type=int)
    per_page = request.args.get("per_page", 50, type=int)

    if per_page not in [25, 50, 100, 200]:
        per_page = 50

    query, filters = build_filtered_query()

    total_filtrado = query.count()

    pagination = query.paginate(
        page=page,
        per_page=per_page,
        error_out=False
    )

    registros = pagination.items

    entry_ids = [registro.id for registro in registros]
    documents_by_entry = {}

    if entry_ids:
        documents = (
            EntryDocument.query
            .filter(EntryDocument.entry_id.in_(entry_ids))
            .all()
        )

        for document in documents:
            documents_by_entry.setdefault(document.entry_id, {})
            documents_by_entry[document.entry_id][document.document_type] = document

    proveedores = (
        Supplier.query
        .filter_by(is_active=True)
        .order_by(Supplier.business_name.asc())
        .all()
    )

    budget_items = get_sorted_budget_items()

    return render_template(
        "registros/index.html",
        registros=registros,
        pagination=pagination,
        proveedores=proveedores,
        budget_items=budget_items,
        filters=filters,
        meses=MESES_ES,
        document_labels=DOCUMENT_LABELS,
        documents_by_entry=documents_by_entry,
        total_filtrado=total_filtrado,
        per_page=per_page
    )


@registros_bp.route("/nuevo", methods=["GET", "POST"])
@login_required
def nuevo():
    proveedores = Supplier.query.filter_by(is_active=True).order_by(Supplier.business_name.asc()).all()
    budget_items = get_sorted_budget_items()

    if request.method == "POST":
        fecha_texto = request.form.get("application_date", "").strip()
        supplier_id = request.form.get("supplier_id", "").strip()
        budget_item_id = request.form.get("budget_item_id", "").strip()
        amount_texto = request.form.get("amount", "").strip()
        operation_folio = request.form.get("operation_folio", "").strip().upper()
        concept = request.form.get("concept", "").strip().upper()
        comments = request.form.get("comments", "").strip().upper()

        application_date = parse_date(fecha_texto)
        amount = parse_amount(amount_texto)

        if not application_date:
            flash("La fecha debe tener formato válido.", "danger")
            return redirect(url_for("registros.nuevo"))

        if not supplier_id or not budget_item_id:
            flash("El proveedor y el rubro son obligatorios.", "danger")
            return redirect(url_for("registros.nuevo"))

        if amount is None or amount <= 0:
            flash("El monto debe ser mayor a cero.", "danger")
            return redirect(url_for("registros.nuevo"))

        if not concept:
            flash("El concepto es obligatorio.", "danger")
            return redirect(url_for("registros.nuevo"))

        month_number = application_date.month
        year = application_date.year
        month_name = MESES_ES[month_number]
        unique_key = generate_unique_key(year)

        registro = Entry(
            unique_key=unique_key,
            application_date=application_date,
            application_month=month_name,
            application_month_number=month_number,
            application_year=year,
            supplier_id=int(supplier_id),
            budget_item_id=int(budget_item_id),
            amount=amount,
            operation_folio=operation_folio or None,
            concept=concept,
            comments=comments or None,
            created_by=current_user.id,
            category_id=None,
            subcategory_id=None
        )

        try:
            db.session.add(registro)
            db.session.flush()

            save_entry_documents_from_request(registro)
            refresh_alerts_after_entry_change(new_budget_item_id=registro.budget_item_id)

            db.session.commit()

            flash(f"Registro {unique_key} creado correctamente.", "success")
            return redirect(url_for("registros.index"))

        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
            return redirect(url_for("registros.nuevo"))

        except Exception as exc:
            db.session.rollback()
            flash(f"No se pudo crear el registro: {exc}", "danger")
            return redirect(url_for("registros.nuevo"))

    return render_template(
        "registros/form.html",
        registro=None,
        proveedores=proveedores,
        budget_items=budget_items,
        document_map={},
        document_labels=DOCUMENT_LABELS,
        modo="nuevo"
    )


@registros_bp.route("/editar/<int:registro_id>", methods=["GET", "POST"])
@login_required
def editar(registro_id):
    registro = Entry.query.get_or_404(registro_id)

    proveedores = Supplier.query.filter_by(is_active=True).order_by(Supplier.business_name.asc()).all()
    budget_items = get_sorted_budget_items()

    if request.method == "POST":
        fecha_texto = request.form.get("application_date", "").strip()
        supplier_id = request.form.get("supplier_id", "").strip()
        budget_item_id = request.form.get("budget_item_id", "").strip()
        amount_texto = request.form.get("amount", "").strip()
        operation_folio = request.form.get("operation_folio", "").strip().upper()
        concept = request.form.get("concept", "").strip().upper()
        comments = request.form.get("comments", "").strip().upper()

        application_date = parse_date(fecha_texto)
        amount = parse_amount(amount_texto)

        if not application_date:
            flash("La fecha debe tener formato válido.", "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

        if not supplier_id or not budget_item_id:
            flash("El proveedor y el rubro son obligatorios.", "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

        if amount is None or amount <= 0:
            flash("El monto debe ser mayor a cero.", "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

        if not concept:
            flash("El concepto es obligatorio.", "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

        month_number = application_date.month
        year = application_date.year
        month_name = MESES_ES[month_number]
        old_budget_item_id = registro.budget_item_id

        try:
            registro.application_date = application_date
            registro.application_month = month_name
            registro.application_month_number = month_number
            registro.application_year = year
            registro.supplier_id = int(supplier_id)
            registro.budget_item_id = int(budget_item_id)
            registro.amount = amount
            registro.operation_folio = operation_folio or None
            registro.concept = concept
            registro.comments = comments or None
            registro.category_id = None
            registro.subcategory_id = None
            registro.updated_at = datetime.utcnow()

            save_entry_documents_from_request(registro)
            refresh_alerts_after_entry_change(
                old_budget_item_id=old_budget_item_id,
                new_budget_item_id=registro.budget_item_id,
            )

            db.session.commit()

            flash("Registro actualizado correctamente.", "success")
            return redirect(url_for("registros.index"))

        except ValueError as exc:
            db.session.rollback()
            flash(str(exc), "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

        except Exception as exc:
            db.session.rollback()
            flash(f"No se pudo actualizar el registro: {exc}", "danger")
            return redirect(url_for("registros.editar", registro_id=registro.id))

    return render_template(
        "registros/form.html",
        registro=registro,
        proveedores=proveedores,
        budget_items=budget_items,
        document_map=get_document_map(registro.id),
        document_labels=DOCUMENT_LABELS,
        modo="editar"
    )


@registros_bp.route("/eliminar/<int:registro_id>", methods=["POST"])
@login_required
def eliminar(registro_id):
    registro = Entry.query.get_or_404(registro_id)

    old_budget_item_id = registro.budget_item_id
    documents = EntryDocument.query.filter_by(entry_id=registro.id).all()

    for document in documents:
        remove_document_file(document)

    db.session.delete(registro)
    db.session.flush()
    refresh_alerts_after_entry_change(old_budget_item_id=old_budget_item_id)
    db.session.commit()

    flash("Registro eliminado correctamente.", "success")
    return redirect(url_for("registros.index"))


@registros_bp.route("/importar-excel", methods=["GET", "POST"])
@login_required
def importar_excel():
    errores = []
    vista_previa = []

    if request.method == "POST":
        archivo = request.files.get("archivo_excel")

        if not archivo or not archivo.filename:
            flash("Selecciona un archivo Excel.", "danger")
            return redirect(url_for("registros.importar_excel"))

        filename = archivo.filename.lower()

        if not filename.endswith((".xlsx", ".xlsm")):
            flash("El archivo debe ser .xlsx o .xlsm.", "danger")
            return redirect(url_for("registros.importar_excel"))

        try:
            wb = load_workbook(archivo, data_only=True)
            ws = wb.worksheets[0]
        except Exception:
            flash("No se pudo leer el Excel. Verifica que el archivo no esté dañado.", "danger")
            return redirect(url_for("registros.importar_excel"))

        headers = get_import_headers(ws)
        missing_headers = [field for field in REQUIRED_IMPORT_FIELDS if field not in headers]

        if missing_headers:
            for field in missing_headers:
                errores.append({
                    "fila": 1,
                    "campo": field.upper(),
                    "valor": "",
                    "detalle": "No se encontró este encabezado en la fila 1."
                })

            return render_template("registros/importar_excel.html", errores=errores, vista_previa=[])

        suppliers = Supplier.query.filter_by(is_active=True).all()
        budget_items = get_sorted_budget_items()
        registros_a_crear = []
        key_counters = {}

        for row in range(2, ws.max_row + 1):
            raw_values = {
                field: ws.cell(row=row, column=headers[field]).value
                for field in REQUIRED_IMPORT_FIELDS
            }

            if all(value is None or str(value).strip() == "" for value in raw_values.values()):
                continue

            application_date = parse_excel_date(raw_values["fecha"])
            amount = parse_excel_amount(raw_values["monto"])
            raw_supplier = raw_values["proveedor"]
            raw_rubro = raw_values["rubro"]
            raw_concept = raw_values["concepto"]

            row_errors = []

            if not application_date:
                row_errors.append({
                    "fila": row,
                    "campo": "FECHA",
                    "valor": raw_values["fecha"],
                    "detalle": "La fecha no tiene un formato válido. Usa dd/mm/aaaa o una fecha real de Excel."
                })

            if amount is None or amount <= 0:
                row_errors.append({
                    "fila": row,
                    "campo": "MONTO",
                    "valor": raw_values["monto"],
                    "detalle": "El monto debe ser mayor a cero."
                })

            supplier, supplier_score = find_supplier_match(raw_supplier, suppliers)

            if not supplier:
                row_errors.append({
                    "fila": row,
                    "campo": "PROVEEDOR",
                    "valor": raw_supplier,
                    "detalle": f"No se encontró un proveedor suficientemente parecido en la base. Mejor coincidencia: {supplier_score:.0%}."
                })

            budget_item = find_budget_item_match(raw_rubro, budget_items)

            if not budget_item:
                row_errors.append({
                    "fila": row,
                    "campo": "RUBRO",
                    "valor": raw_rubro,
                    "detalle": "No existe un rubro activo con ese número/código."
                })

            concept = normalize_text(raw_concept)

            if not concept:
                row_errors.append({
                    "fila": row,
                    "campo": "CONCEPTO",
                    "valor": raw_concept,
                    "detalle": "El concepto es obligatorio."
                })

            if row_errors:
                errores.extend(row_errors)
                continue

            month_number = application_date.month
            year = application_date.year
            month_name = MESES_ES[month_number]
            unique_key = generate_unique_key_for_import(year, key_counters)

            registro = Entry(
                unique_key=unique_key,
                application_date=application_date,
                application_month=month_name,
                application_month_number=month_number,
                application_year=year,
                supplier_id=supplier.id,
                budget_item_id=budget_item.id,
                amount=amount,
                operation_folio=normalize_text(raw_values["folio"]) or None,
                concept=concept,
                comments=normalize_text(raw_values["comentarios"]) or None,
                created_by=current_user.id,
                category_id=None,
                subcategory_id=None
            )

            registros_a_crear.append(registro)

            vista_previa.append({
                "unique_key": unique_key,
                "fecha": application_date.strftime("%d/%m/%Y"),
                "proveedor": supplier.business_name,
                "rubro": f"{budget_item.item_code} - {budget_item.item_name}",
                "monto": f"{amount:,.2f}",
            })

        if errores:
            db.session.rollback()
            flash("El Excel fue rechazado. Corrige los proveedores, rubros o campos marcados.", "danger")
            return render_template("registros/importar_excel.html", errores=errores, vista_previa=[])

        if not registros_a_crear:
            flash("El Excel no contiene filas con información para importar.", "warning")
            return redirect(url_for("registros.importar_excel"))

        db.session.add_all(registros_a_crear)
        db.session.flush()
        refresh_alerts_for_budget_item_ids({registro.budget_item_id for registro in registros_a_crear})
        db.session.commit()

        flash(f"Se importaron {len(registros_a_crear)} registros correctamente.", "success")
        return render_template("registros/importar_excel.html", errores=[], vista_previa=vista_previa)

    return render_template("registros/importar_excel.html", errores=[], vista_previa=[])


# =========================================================
# RUTAS DE DOCUMENTOS
# =========================================================

@registros_bp.route("/documento/<int:document_id>/ver")
@login_required
def ver_documento(document_id):
    document = EntryDocument.query.get_or_404(document_id)

    if not document.file_path or not os.path.exists(document.file_path):
        abort(404)

    extension = (document.file_extension or "").lower()

    if extension not in PREVIEWABLE_EXTENSIONS:
        return redirect(url_for("registros.descargar_documento", document_id=document.id))

    if extension == "pdf":
        mimetype = "application/pdf"
    elif extension == "png":
        mimetype = "image/png"
    elif extension in ["jpg", "jpeg"]:
        mimetype = "image/jpeg"
    else:
        mimetype = "application/octet-stream"

    response = send_file(
        document.file_path,
        mimetype=mimetype,
        as_attachment=False,
        download_name=document.original_filename,
        conditional=False
    )

    response.headers["Content-Disposition"] = f'inline; filename="{document.original_filename}"'
    response.headers["X-Content-Type-Options"] = "nosniff"

    return response


@registros_bp.route("/documento/<int:document_id>/descargar")
@login_required
def descargar_documento(document_id):
    document = EntryDocument.query.get_or_404(document_id)

    if not document.file_path or not os.path.exists(document.file_path):
        abort(404)

    return send_file(
        document.file_path,
        as_attachment=True,
        download_name=document.original_filename
    )


@registros_bp.route("/documento/<int:document_id>/eliminar", methods=["POST"])
@login_required
def eliminar_documento(document_id):
    document = EntryDocument.query.get_or_404(document_id)
    registro_id = document.entry_id

    remove_document_file(document)

    db.session.delete(document)
    db.session.commit()

    flash("Documento eliminado correctamente.", "success")
    return redirect(url_for("registros.editar", registro_id=registro_id))


@registros_bp.route("/<int:registro_id>/descargar-documentos")
@login_required
def descargar_documentos_registro(registro_id):
    registro = Entry.query.get_or_404(registro_id)

    documents = (
        EntryDocument.query
        .filter_by(entry_id=registro.id)
        .order_by(EntryDocument.document_type.asc())
        .all()
    )

    if not documents:
        flash("Este registro todavía no tiene documentos cargados.", "warning")
        return redirect(url_for("registros.index"))

    memory_file = BytesIO()

    with zipfile.ZipFile(memory_file, "w", zipfile.ZIP_DEFLATED) as zf:
        for document in documents:
            if not document.file_path or not os.path.exists(document.file_path):
                continue

            label = DOCUMENT_LABELS.get(document.document_type, document.document_type)
            safe_label = secure_filename(label.replace(" ", "_").lower())
            extension = document.file_extension or "file"
            arcname = f"{registro.unique_key}/{safe_label}_{document.id}.{extension}"

            zf.write(document.file_path, arcname)

    memory_file.seek(0)

    return send_file(
        memory_file,
        as_attachment=True,
        download_name=f"evidencias_{registro.unique_key}.zip",
        mimetype="application/zip"
    )


@registros_bp.route("/descargar-documentos-filtrados")
@login_required
def descargar_documentos_filtrados():
    query, filters = build_filtered_query()
    registros = query.all()

    if not registros:
        flash("No hay registros con los filtros actuales.", "warning")
        return redirect(url_for("registros.index"))

    entry_ids = [registro.id for registro in registros]
    registros_por_id = {registro.id: registro for registro in registros}

    documents = (
        EntryDocument.query
        .filter(EntryDocument.entry_id.in_(entry_ids))
        .order_by(EntryDocument.entry_id.asc(), EntryDocument.document_type.asc())
        .all()
    )

    if not documents:
        flash("Los registros filtrados no tienen documentos cargados.", "warning")
        return redirect(url_for("registros.index", **filters))

    memory_file = BytesIO()

    with zipfile.ZipFile(memory_file, "w", zipfile.ZIP_DEFLATED) as zf:
        for document in documents:
            if not document.file_path or not os.path.exists(document.file_path):
                continue

            registro = registros_por_id.get(document.entry_id)

            if not registro:
                continue

            label = DOCUMENT_LABELS.get(document.document_type, document.document_type)
            safe_label = secure_filename(label.replace(" ", "_").lower())
            extension = document.file_extension or "file"

            arcname = f"{registro.unique_key}/{safe_label}_{document.id}.{extension}"
            zf.write(document.file_path, arcname)

    memory_file.seek(0)

    return send_file(
        memory_file,
        as_attachment=True,
        download_name="evidencias_registros_filtrados.zip",
        mimetype="application/zip"
    )